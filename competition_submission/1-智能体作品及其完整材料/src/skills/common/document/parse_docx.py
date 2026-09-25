import base64
import io
from typing import Any


def execute(payload: dict[str, Any]) -> dict[str, Any]:
    paragraphs: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    warnings: list[str] = []
    if payload.get("content_base64"):
        try:
            from docx import Document

            document = Document(io.BytesIO(base64.b64decode(payload["content_base64"])))
            for index, paragraph in enumerate(document.paragraphs, start=1):
                text = paragraph.text.strip()
                if text:
                    paragraphs.append({"paragraph": index, "text": text, "style": getattr(paragraph.style, "name", None)})
            for table_index, table in enumerate(document.tables, start=1):
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                tables.append({"table": table_index, "rows": rows})
        except Exception as exc:
            warnings.append(f"docx parse failed: {exc}")
    elif payload.get("text"):
        for index, line in enumerate(payload["text"].splitlines(), start=1):
            if line.strip():
                paragraphs.append({"paragraph": index, "text": line.strip(), "style": None})
        warnings.append("no docx bytes supplied; using provided text as paragraphs")
    else:
        warnings.append("no docx content supplied")
    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "text": "\n".join(item["text"] for item in paragraphs),
        "warnings": warnings,
    }
