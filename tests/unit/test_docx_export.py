from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone

import pytest
from docx import Document
from fastapi.testclient import TestClient

from engine import docx_export
from engine.api import app
from engine.settings import settings

CASE = "11111111-1111-1111-1111-111111111111"
ARTIFACT = "22222222-2222-2222-2222-222222222222"
FACTS = "33333333-3333-3333-3333-333333333333"


def payload(body="第一行\t含制表符\n第二行\n中文"):
    return {
        "case_id": CASE,
        "artifact_version_id": ARTIFACT,
        "version": 2,
        "schema_version": "draft.v2",
        "doc_type": "起诉意见书",
        "body": body,
        "created_at": datetime(2026, 1, 2, 3, 4, 5, tzinfo=timezone.utc),
        "facts_version_id": FACTS,
    }


def test_real_docx_is_deterministic_and_contains_provenance():
    first = docx_export.render_docx(payload())
    second = docx_export.render_docx(payload())
    assert first == second
    with zipfile.ZipFile(io.BytesIO(first)) as archive:
        assert "word/document.xml" in archive.namelist()
        assert "第一行" in archive.read("word/document.xml").decode()
        styles = archive.read("word/styles.xml").decode()
        title_style = re.search(r'<w:style[^>]*w:styleId="Title".*?</w:style>', styles)
        assert title_style is not None and "w:pBdr" not in title_style.group(0)
        core = archive.read("docProps/core.xml").decode()
        assert "2026-01-02T03:04:05Z" in core


def test_body_line_breaks_and_trailing_blank_line_are_preserved():
    body = "第一行\r\n\r\n末行\n"
    document = Document(io.BytesIO(docx_export.render_docx(payload(body))))
    body_paragraphs = [paragraph.text for paragraph in document.paragraphs[2:-1]]
    assert "\n".join(body_paragraphs) == body.replace("\r\n", "\n").replace("\r", "\n")


@pytest.mark.parametrize("body", ["等待【待核实】", "等待{{facts.missing}}"])
def test_unresolved_placeholders_are_rejected(body):
    with pytest.raises(docx_export.DocxExportError) as error:
        docx_export.render_docx(payload(body))
    assert error.value.code == "DRAFT_EXPORT_BLOCKED"


def test_xml_control_character_is_rejected():
    with pytest.raises(docx_export.DocxExportError) as error:
        docx_export.render_docx(payload("有效\x00内容"))
    assert error.value.code == "INVALID_INPUT"


@pytest.mark.parametrize("body", ["坏\ud800字符", "坏\uffff字符"])
def test_invalid_xml_unicode_is_rejected(body):
    with pytest.raises(docx_export.DocxExportError) as error:
        docx_export.render_docx(payload(body))
    assert error.value.code == "INVALID_INPUT"


def test_naive_created_at_is_rejected_for_determinism():
    value = payload()
    value["created_at"] = datetime(2026, 1, 2, 3, 4, 5)
    with pytest.raises(docx_export.DocxExportError) as error:
        docx_export.render_docx(value)
    assert error.value.code == "INVALID_INPUT"


def test_internal_route_requires_token_and_returns_docx(monkeypatch):
    monkeypatch.setattr(settings, "service_token", "test-service-token")
    client = TestClient(app)
    assert client.post("/internal/v1/draft-exports/docx", json={}).status_code == 401
    response = client.post(
        "/internal/v1/draft-exports/docx",
        headers={"X-Service-Token": "test-service-token"},
        json={**payload(), "created_at": "2026-01-02T03:04:05Z"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(docx_export.DOCX_MIME)
    assert response.content.startswith(b"PK")


def test_internal_route_rejects_naive_date_and_unknown_fields(monkeypatch):
    monkeypatch.setattr(settings, "service_token", "test-service-token")
    client = TestClient(app)
    headers = {"X-Service-Token": "test-service-token"}
    naive = client.post(
        "/internal/v1/draft-exports/docx", headers=headers,
        json={**payload(), "created_at": "2026-01-02T03:04:05"},
    )
    assert naive.status_code == 422
    extra = client.post(
        "/internal/v1/draft-exports/docx", headers=headers,
        json={**payload(), "created_at": "2026-01-02T03:04:05Z", "storage_key": "client-controlled"},
    )
    assert extra.status_code == 422
