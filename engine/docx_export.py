"""Deterministic Word export for immutable draft artifacts.

The application service owns authorization and persistence.  This module only
accepts the already validated immutable artifact envelope and renders a real
editable DOCX document.  Keeping the renderer here avoids exposing the engine
or object storage to browsers.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone
from typing import Any, Literal
from uuid import UUID

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field, StrictInt, field_validator

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_INVALID_XML_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_PLACEHOLDER = re.compile(r"\{\{\s*[^{}]+?\s*\}\}")
_CHINESE_PLACEHOLDER = re.compile(r"【[^】]*】")


class DocxExportError(ValueError):
    """A client supplied export envelope cannot be rendered safely."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


class DocxExportRequest(BaseModel):
    """Strict internal contract; clients cannot smuggle storage or status fields."""

    model_config = ConfigDict(extra="forbid")

    case_id: UUID
    artifact_version_id: UUID
    version: StrictInt = Field(ge=1)
    schema_version: Literal["draft.v2", "case.draft.v1"]
    doc_type: str
    body: str
    created_at: datetime
    facts_version_id: UUID | None = None

    @field_validator("created_at")
    @classmethod
    def require_timezone(cls, value: datetime) -> datetime:
        if value.tzinfo is None or value.utcoffset() is None:
            raise ValueError("created_at must include a timezone")
        return value


def _require_uuid(value: Any, field: str) -> str:
    try:
        return str(UUID(str(value)))
    except (ValueError, TypeError, AttributeError) as exc:
        raise DocxExportError("INVALID_INPUT", f"{field} must be a UUID") from exc


def _validate_text(value: Any, field: str, *, nonblank: bool = False) -> str:
    if not isinstance(value, str):
        raise DocxExportError("INVALID_INPUT", f"{field} must be text")
    if nonblank and not value.strip():
        raise DocxExportError("DRAFT_EXPORT_BLOCKED", f"{field} is blank")
    if _INVALID_XML_CONTROL.search(value) or any(
        0xD800 <= ord(char) <= 0xDFFF
        or 0xFDD0 <= ord(char) <= 0xFDEF
        or ord(char) in {0xFFFE, 0xFFFF}
        or ord(char) >= 0x10000 and ord(char) % 0x10000 in {0xFFFE, 0xFFFF}
        for char in value
    ):
        raise DocxExportError("INVALID_INPUT", f"{field} contains invalid XML control characters")
    return value


def _set_font(run, *, size: int = 12, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Noto Sans CJK SC"
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    east_asia = rpr.rFonts
    if east_asia is None:
        east_asia = OxmlElement("w:rFonts")
        rpr.insert(0, east_asia)
    east_asia.set(qn("w:eastAsia"), "Noto Sans CJK SC")


def _set_cell_or_paragraph_font(style) -> None:
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")


def _remove_paragraph_borders(element) -> None:
    for border in list(element.iter(qn("w:pBdr"))):
        parent = border.getparent()
        if parent is not None:
            parent.remove(border)


def _normalise_zip(raw: bytes) -> bytes:
    """Remove ZIP timestamps/order variability introduced by python-docx."""
    source = io.BytesIO(raw)
    output = io.BytesIO()
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as zout:
        for name in sorted(zin.namelist()):
            info = zin.getinfo(name)
            data = zin.read(name)
            target = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            target.compress_type = zipfile.ZIP_DEFLATED
            target.external_attr = info.external_attr
            target.create_system = info.create_system
            target.flag_bits = info.flag_bits
            zout.writestr(target, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    return output.getvalue()


def render_docx(payload: dict[str, Any]) -> bytes:
    """Render an immutable artifact envelope into deterministic DOCX bytes."""
    if not isinstance(payload, dict):
        raise DocxExportError("INVALID_INPUT", "payload must be an object")
    schema = payload.get("schema_version")
    if schema not in {"draft.v2", "case.draft.v1"}:
        raise DocxExportError("INVALID_INPUT", "unsupported schema_version")
    artifact_id = _require_uuid(payload.get("artifact_version_id"), "artifact_version_id")
    _require_uuid(payload.get("case_id"), "case_id")
    facts_id = payload.get("facts_version_id")
    facts_id = _require_uuid(facts_id, "facts_version_id") if facts_id is not None else None
    version = payload.get("version")
    if not isinstance(version, int) or isinstance(version, bool) or version < 1:
        raise DocxExportError("INVALID_INPUT", "version must be a positive integer")
    doc_type = _validate_text(payload.get("doc_type"), "doc_type", nonblank=True)
    body = _validate_text(payload.get("body"), "body", nonblank=True)
    if _PLACEHOLDER.search(body) or _CHINESE_PLACEHOLDER.search(body):
        raise DocxExportError("DRAFT_EXPORT_BLOCKED", "body contains unresolved placeholders")
    created_at = payload.get("created_at")
    if not isinstance(created_at, datetime) or created_at.tzinfo is None or created_at.utcoffset() is None:
        raise DocxExportError("INVALID_INPUT", "created_at must be an ISO timestamp")
    created_at = created_at.astimezone(timezone.utc).replace(tzinfo=None)

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    _set_cell_or_paragraph_font(document.styles["Normal"])

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("文书辅助稿，不替代司法裁量")
    _set_font(header_run, size=10, color=RGBColor(0, 0, 0))

    title = document.add_paragraph()
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    _remove_paragraph_borders(title_style.element)
    title.style = title_style
    _remove_paragraph_borders(title._p)
    title.paragraph_format.keep_with_next = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(doc_type)
    _set_font(title_run, size=15, bold=True, color=RGBColor(0, 0, 0))
    label = document.add_paragraph()
    label.paragraph_format.keep_with_next = True
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_run = label.add_run("文书辅助稿，不替代司法裁量")
    _set_font(label_run, size=10, color=RGBColor(90, 90, 90))

    for line in body.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(line)
        _set_font(run)

    provenance = document.add_paragraph()
    provenance.paragraph_format.space_before = Pt(14)
    provenance_run = provenance.add_run(
        f"来源工件版本：{artifact_id}（版本 {version}）"
        + (f"；事实版本：{facts_id}" if facts_id else "")
    )
    _set_font(provenance_run, size=9, color=RGBColor(100, 100, 100))

    core = document.core_properties
    core.author = "LexCyber"
    core.title = f"{doc_type} 文书辅助稿"
    core.subject = "LexCyber immutable draft export"
    core.comments = f"artifact_version_id={artifact_id};version={version}"
    core.created = created_at
    core.modified = created_at
    core.last_printed = created_at

    raw = io.BytesIO()
    document.save(raw)
    return _normalise_zip(raw.getvalue())
